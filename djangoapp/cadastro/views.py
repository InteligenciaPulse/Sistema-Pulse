import os
import json
import logging
import tempfile
import openpyxl
from django.conf import settings
from weasyprint import HTML, CSS
from django.db import transaction
from openpyxl.styles import numbers
from django.http import JsonResponse
from django.http import HttpResponse
from django.utils.timezone import now
from django.middleware.csrf import get_token
from django.shortcuts import render, redirect
from django.shortcuts import get_object_or_404
from django.views.decorators.csrf import csrf_protect
from .models import SolicitacaoOrcamento, Orcamento, Paciente, Procedimento, Produto, Parceiro, ParceiroProdutos, Subtipo, Pacote, PacoteProcedimentos, Endereco, Status, OrcamentoParceiros, Especialidade, Custos, OrcamentoProcedimentos

from .models import Coluna

from django.contrib.auth import authenticate, login
from django.http import JsonResponse
from rest_framework.decorators import api_view

@api_view(['POST'])
def login_view(request):
    username = request.data.get('username')
    password = request.data.get('password')
    user = authenticate(request, username=username, password=password)
    if user is not None:
        login(request, user)
        return JsonResponse({"success": True})
    else:
        return JsonResponse({"error": "Invalid credentials"}, status=400)
    

    

def home(request):
    return render(request, 'cadastro/home.html')

def criar_orcamento(request):
    parceiro_list = Parceiro.objects.all().order_by('nome')
    status_list = Status.objects.all()
    especialidade_list = Especialidade.objects.all().order_by('nome')
    subtipo_list = Subtipo.objects.all().order_by('nome')
    procedimento_list = Procedimento.objects.all().order_by('nome')

    context = {
        "parceiro_list": parceiro_list,
        "status_list": status_list,
        "especialidade_list": especialidade_list,
        "subtipo_list": subtipo_list,
        'procedimento_list': procedimento_list
    }

    return render(request, 'cadastro/criar_orcamento.html', context)

def historico(request):
    tipo = request.GET.get("tipo", "orcamento")

    status_id = request.GET.get("status")
    paciente_nome = request.GET.get("paciente")
    especialidade_id = request.GET.get("especialidade")
    procedimento_nome = request.GET.get("procedimento")
    parceiro_nome = request.GET.get("parceiro")
    data_criacao = request.GET.get("data_criacao")
    data_aprovacao = request.GET.get("data_aprovacao")

    if tipo == "solicitacoes":
        atividades = SolicitacaoOrcamento.objects.select_related("paciente", "orcamento", "status").order_by("-data_solicitacao")
        
        if status_id:
            atividades = atividades.filter(status_id=status_id)
        
        if paciente_nome:
            atividades = atividades.filter(paciente__nome__icontains=paciente_nome)
        
        if especialidade_id:
            atividades = atividades.filter(especialidade_id=especialidade_id)
        
        if data_criacao:
            atividades = atividades.filter(data_solicitacao=data_criacao)

    else:
        atividades = Orcamento.objects.select_related("status").prefetch_related(
            "solicitacao_orcamento__paciente",
            "orcamento_parceiros__parceiro",
            "orcamento_parceiros__produto"
        ).order_by("-data_criacao")

        if status_id:
            atividades = atividades.filter(status_id=status_id)
        
        if paciente_nome:
            atividades = atividades.filter(solicitacao_orcamento__paciente__nome__icontains=paciente_nome)
        
        if especialidade_id:
            atividades = atividades.filter(solicitacao_orcamento__especialidade_id=especialidade_id)
        
        if procedimento_nome:
            atividades = atividades.filter(orcamento_parceiros__produto__nome__icontains=procedimento_nome)
        
        if parceiro_nome:
            atividades = atividades.filter(orcamento_parceiros__parceiro__nome__icontains=parceiro_nome)
        
        if data_criacao:
            atividades = atividades.filter(data_criacao=data_criacao)
        
        if data_aprovacao:
            atividades = atividades.filter(data_aprovacao=data_aprovacao)

    status_list = Status.objects.all()
    especialidade_list = Especialidade.objects.all()

    context = {
        "atividades": atividades,
        "tipo": tipo,
        "status_list": status_list,
        "especialidade_list": especialidade_list,
        "filtros": {
            "status": status_id,
            "paciente": paciente_nome,
            "especialidade": especialidade_id,
            "procedimento": procedimento_nome,
            "parceiro": parceiro_nome,
            "data_criacao": data_criacao,
            "data_aprovacao": data_aprovacao,
        }
    }

    return render(request, "cadastro/historico.html", context)

def buscar_pacientes(request):
    query = request.GET.get('q', '')
    pacientes = Paciente.objects.filter(nome__icontains=query)[:10]
    data = [{"id": p.id, "nome": p.nome} for p in pacientes]
    return JsonResponse(data, safe=False)

def buscar_procedimentos(request):
    query = request.GET.get('q', '').strip()
    
    if not query:
        return JsonResponse([], safe=False)
    
    procedimentos = Procedimento.objects.filter(nome__icontains=query).values("id", "nome")
    return JsonResponse(list(procedimentos), safe=False)

def buscar_parceiros_by(request):
    query = request.GET.get('q', '').strip()
    subtipo_id = request.GET.get('subtipo', '').strip()
    
    if not query:
        return JsonResponse([], safe=False)
    
    parceiros = Parceiro.objects.filter(nome__icontains=query)

    if subtipo_id:
        parceiros = parceiros.filter(subtipo_id=subtipo_id)
    
    parceiros = parceiros.values("id", "nome")
    
    return JsonResponse(list(parceiros), safe=False)

def buscar_procedimentos_by(request):
    query = request.GET.get('q', '').strip()
    especialidade_id = request.GET.get('especialidade', '').strip()
    
    if not query:
        return JsonResponse([], safe=False)
    
    procedimentos = Procedimento.objects.filter(nome__icontains=query)

    if especialidade_id:
        procedimentos = procedimentos.filter(especialidade_id=especialidade_id)
    
    procedimentos = procedimentos.values("id", "nome")
    
    return JsonResponse(list(procedimentos), safe=False)

def buscar_parceiros(request):
    query = request.GET.get('q', '')
    parceiros = Parceiro.objects.filter(nome__icontains=query)[:10]
    data = [{"id": p.id, "nome": p.nome} for p in parceiros]
    return JsonResponse(data, safe=False)

# def buscar_parceiros_por_procedimento(request):
#     procedimento_nome = request.GET.get("procedimento_nome", "").strip()

#     if not procedimento_nome:
#         return JsonResponse([], safe=False)

#     try:
#         procedimento = Procedimento.objects.get(nome=procedimento_nome)
#         procedimento_id = procedimento.id

#         parceiros_procedimentos = ParceiroProcedimentos.objects.filter(procedimento_id=procedimento_id)

#         resposta = [
#             {
#                 "id": pp.parceiro.id,
#                 "nome": pp.parceiro.nome,
#                 "valor_venda": pp.valor_venda,
#             }
#             for pp in parceiros_procedimentos
#         ]

#         return JsonResponse(resposta, safe=False)

#     except Procedimento.DoesNotExist:
#         return JsonResponse([], safe=False)

def buscar_produtos_por_parceiro(request):
    parceiro_nome = request.GET.get("parceiro_nome", "").strip()

    if not parceiro_nome:
        return JsonResponse([], safe=False)

    try:
        parceiro = Parceiro.objects.get(nome=parceiro_nome)
        parceiro_id = parceiro.id

        produtos_parceiro = ParceiroProdutos.objects.filter(parceiro_id=parceiro_id)

        resposta = [
            {
                "id": pp.produto.id,
                "nome": pp.produto.nome,
                "valor_venda": pp.valor_venda,
                "valor_particular": pp.valor_particular,
                "valor_repasse": pp.valor_repasse
            }
            for pp in produtos_parceiro
        ]

        return JsonResponse(resposta, safe=False)

    except Parceiro.DoesNotExist:
        return JsonResponse([], safe=False)

def buscar_pacotes(request):
    query = request.GET.get('q', '').strip()
    
    if not query:
        return JsonResponse([], safe=False)
    
    pacotes = Pacote.objects.filter(nome__icontains=query).values("id", "nome")
    return JsonResponse(list(pacotes), safe=False)

def buscar_procedimentos_por_pacote(request):
    pacote_nome = request.GET.get("pacote_nome", "").strip()

    try:
        pacote = Pacote.objects.get(nome=pacote_nome)

        procedimentos_relacionados = PacoteProcedimentos.objects.filter(pacote=pacote)
        
        procedimentos = [
            {"id": p.procedimento.id, "nome": p.procedimento.nome}
            for p in procedimentos_relacionados
        ]

        return JsonResponse(procedimentos, safe=False)
    except Pacote.DoesNotExist:
        return JsonResponse([], safe=False)
    
def buscar_subtipos(request):
    subtipos = Subtipo.objects.all().values("id", "nome")
    return JsonResponse(list(subtipos), safe=False)

def buscar_especialidades(request):
    especialidades = Especialidade.objects.all().values("id", "nome")
    return JsonResponse(list(especialidades), safe=False)

def buscar_parceiros_por_subtipo(request):
    subtipo_id = request.GET.get("subtipo")
    parceiros = ParceiroProdutos.objects.filter(parceiro__subtipo_id=subtipo_id).values("parceiro__id", "parceiro__nome", "valor_venda")
    return JsonResponse(list(parceiros), safe=False)

def visualizar_orcamento_pdf(request):
    cliente = {
        "nome": "Maria Estetiane da Silva",
        "cpf": "000.000.00-00",
        "telefone": "(87) 9 8176-0222",
        "email": "gestão@meupulse.com.br",
        "endereco": "R. Dr. Júlio de Melo, 538 - Centro, Petrolina - PE"
    }

    procedimentos = [
        {"descricao": "Procedimento X", "observacao": "Observação Y"},
        {"descricao": "Procedimento Z", "observacao": "Observação W"},
    ]

    valor_total = "00.000,00"

    # Renderiza o HTML
    html_string = render(request, "cadastro/orcamento.html", {
        "cliente": cliente,
        "procedimentos": procedimentos,
        "valor_total": valor_total
    }).content.decode("utf-8")

    # Caminho absoluto do CSS estático
    css_path = os.path.join(settings.BASE_DIR, 'cadastro/static/cadastro/css/orcamento_style.css')

    # Geração do PDF com WeasyPrint usando o CSS externo
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'inline; filename="orcamento.pdf"'
    
    with tempfile.NamedTemporaryFile(delete=True) as temp_file:
        HTML(string=html_string).write_pdf(temp_file.name, stylesheets=[CSS(css_path)])
        response.write(open(temp_file.name, "rb").read())

    return response

def visualizar_orcamento_html(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)

    try:
        solicitacao = SolicitacaoOrcamento.objects.get(orcamento=orcamento)
    except SolicitacaoOrcamento.DoesNotExist:
        return JsonResponse({"error": "Solicitação de orçamento não encontrada."}, status=404)
    
    paciente = solicitacao.paciente
    produtos = OrcamentoParceiros.objects.filter(orcamento=orcamento).select_related("produto", "parceiro")
    
    pp = OrcamentoParceiros.objects.filter(orcamento=orcamento)
    print('fff', pp)

    cliente = {
        "nome": paciente.nome,
        "cpf": paciente.cpf,
        "telefone": paciente.telefone,
        "email": paciente.email,
        "endereco": f"{paciente.endereco.rua}, {paciente.endereco.numero} - {paciente.endereco.bairro}, {paciente.endereco.cidade} - {paciente.endereco.estado}, {paciente.endereco.cep}",
    }

    produtos_lista = [
        {
            "descricao": f"{prod.produto.nome} - {prod.parceiro.nome}",
            "observacao": f"Valor: R$ {prod.valor_venda:.2f} - Repasse: R$ {prod.valor_repasse:.2f}"
        }
        for prod in produtos
    ]
    
    return render(request, "cadastro/orcamento.html", {
        "cliente": cliente,
        "produtos": produtos_lista,
        "valor_total": f"R$ {orcamento.valor_total:.2f}"
    })

@csrf_protect
def salvar_paciente(request):
    if request.method == 'POST':
        try:
            with transaction.atomic():
                # Capturando dados do endereço
                rua = request.POST.get('rua')
                numero = request.POST.get('numero')
                bairro = request.POST.get('bairro')
                cidade = request.POST.get('cidade')
                estado = request.POST.get('estado')
                cep = request.POST.get('cep')
                complemento = request.POST.get('complemento', '')
                referencia = request.POST.get('referencia', '')

                # Criando e salvando endereço
                endereco = Endereco.objects.create(
                    rua=rua, numero=numero, bairro=bairro, cidade=cidade,
                    estado=estado, cep=cep, complemento=complemento, referencia=referencia
                )

                # Capturando dados do paciente
                nome = request.POST.get('nome')
                cpf = request.POST.get('cpf')
                telefone = request.POST.get('telefone')
                email = request.POST.get('email', '')
                data_nascimento = request.POST.get('data_nascimento')
                genero = request.POST.get('genero')
                anamnese = request.POST.get('anamnese', '')

                # Criando e salvando paciente
                paciente = Paciente.objects.create(
                    nome=nome, cpf=cpf, telefone=telefone, email=email,
                    data_nascimento=data_nascimento, genero=genero, anamnese=anamnese, endereco=endereco
                )

            return JsonResponse({'message': 'Paciente salvo com sucesso!'}, status=201)

        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)

    return JsonResponse({'error': 'Método não permitido'}, status=405)

@csrf_protect
def salvar_orcamento(request):
    if request.method == "POST":
        try:
            data = json.loads(request.body)

            with transaction.atomic():
                paciente_id = data.get("paciente_id")
                paciente = Paciente.objects.get(id=paciente_id)
                status = Status.objects.get(id=data.get('status'))

                orcamento = Orcamento.objects.create(
                    status=status,
                    valor_total=data.get("valor_total"),
                    data_criacao=now()
                )

                solicitacao = SolicitacaoOrcamento.objects.create(
                    paciente=paciente,
                    orcamento=orcamento,
                    status=status,
                    data_solicitacao=now()
                )

                for item in data.get('procedimentos', []):
                    procedimento_id = item['procedimento_id']

                    OrcamentoProcedimentos.objects.create(
                        orcamento=orcamento,
                        procedimento_id=procedimento_id
                    )

                for item in data.get("produtos", []):
                    parceiro_id = item["parceiro_id"]
                    produto_id = item["produto_id"]
                    valor_venda = item["valor_venda"]

                    parceiro_produto = ParceiroProdutos.objects.get(
                        parceiro_id=parceiro_id, 
                        produto_id=produto_id
                    )

                    custos = Custos.objects.create(
                        comissao = item['comissao'],
                        brindes = item['brindes'],
                        imposto = item['impostos'],
                        cartao = item['cartoes']
                    )

                    OrcamentoParceiros.objects.create(
                        orcamento=orcamento,
                        parceiro_id=parceiro_id,
                        produto_id=produto_id,
                        valor_venda=valor_venda,
                        valor_repasse=parceiro_produto.valor_repasse,
                        custos=custos,
                        margem_lucro = item['margem_lucro']
                    )

            return JsonResponse({"success": True, "message": "Orçamento salvo com sucesso!", "orcamento_id": orcamento.id}, status=201)

        except Paciente.DoesNotExist:
            return JsonResponse({"error": "Paciente não encontrado."}, status=404)
        except ParceiroProdutos.DoesNotExist:
            return JsonResponse({"error": "Parceiro ou Produto inválido."}, status=400)
        except Exception as e:
            return JsonResponse({"error": str(e)}, status=500)

    return JsonResponse({"error": "Método não permitido."}, status=405)

def atualizar_status(request, orcamento_id):
    if request.method == "POST":
        try:
            data = json.loads(request.body)
            novo_status = data.get("status")

            orcamento = Orcamento.objects.get(id=orcamento_id)
            orcamento.status_id = novo_status
            orcamento.save()

            return JsonResponse({"success": True})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)}, status=400)
    return JsonResponse({"error": "Método não permitido"}, status=405)

def editar_orcamento(request, orcamento_id):
    parceiro_list = Parceiro.objects.all().order_by('nome')
    status_list = Status.objects.all()
    especialidade_list = Especialidade.objects.all().order_by('nome')
    subtipo_list = Subtipo.objects.all().order_by('nome')
    procedimento_list = Procedimento.objects.all().order_by('nome')

    orcamento = get_object_or_404(Orcamento, id=orcamento_id)

    try:
        solicitacao = SolicitacaoOrcamento.objects.get(orcamento=orcamento)
        paciente = solicitacao.paciente
    except SolicitacaoOrcamento.DoesNotExist:
        paciente = None
        
    parceiros_dict = {}
    orcamento_parceiros = OrcamentoParceiros.objects.filter(orcamento=orcamento)
    total_particular = 0

    for orc_parc in orcamento_parceiros:
        parceiro_id = orc_parc.parceiro.id
        parceiro_produto = ParceiroProdutos.objects.get(parceiro_id=parceiro_id, produto_id=orc_parc.produto.id)

        if parceiro_id not in parceiros_dict:
            parceiros_dict[parceiro_id] = {
                "parceiro_id": parceiro_id,
                "parceiro_nome": orc_parc.parceiro.nome,
                "produtos": []
            }
        
        parceiros_dict[parceiro_id]["produtos"].append({
            "produto_id": orc_parc.produto.id,
            "produto_nome": orc_parc.produto.nome,
            "valor_venda": str(orc_parc.valor_venda),
            "valor_repasse": str(orc_parc.valor_repasse),
            "valor_particular": parceiro_produto.valor_particular
        })

        total_particular += parceiro_produto.valor_particular

    procedimento_dict = {}
    orc_procedimentos = OrcamentoProcedimentos.objects.filter(orcamento=orcamento)

    for orc_proced in orc_procedimentos:
        procedimento_id = orc_proced.procedimento.id

        if procedimento_id not in procedimento_dict:
            procedimento_dict[procedimento_id] = {
                "procedimento_id": procedimento_id,
                "procedimento_nome": orc_proced.procedimento.nome
            }

    context = {
        "orcamento": orcamento,
        "orc_status": orcamento.status,
        "paciente": paciente,
        "parceiros": list(parceiros_dict.values()),
        "parceiro_list": parceiro_list,
        "status_list": status_list,
        "especialidade_list": especialidade_list,
        "subtipo_list": subtipo_list,
        'orc_procedimento': list(procedimento_dict.values()),
        'procedimento_list': procedimento_list,
        'total_particular': f"{total_particular:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    }
    
    return render(request, "cadastro/editar_orcamento.html", context)

@csrf_protect
def atualizar_orcamento(request, orcamento_id=None):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            orcamento_id = data.get("orcamento_id", orcamento_id)
            status = Status.objects.get(id=data.get('status'))

            if not orcamento_id:
                return JsonResponse({"error": "Orçamento ID não fornecido"}, status=400)

            orcamento = Orcamento.objects.get(id=orcamento_id)

            with transaction.atomic():
                orcamento.valor_total = data["valor_total"]
                orcamento.status = status
                orcamento.save()
                
                OrcamentoParceiros.objects.filter(orcamento=orcamento).delete()

                for prod_data in data["produtos"]:
                    parceiro = Parceiro.objects.get(id=prod_data["parceiro_id"])
                    produto = Produto.objects.get(id=prod_data["produto_id"])

                    orc = OrcamentoParceiros.objects.create(
                        orcamento=orcamento,
                        parceiro=parceiro,
                        produto=produto,
                        valor_venda=prod_data["valor_venda"],
                        valor_repasse=ParceiroProdutos.objects.get(parceiro=parceiro, produto=produto).valor_repasse
                    )
            return JsonResponse({"message": "Orçamento atualizado com sucesso!"})

        except Exception as e:
            return JsonResponse({"error": str(e)}, status=400)

    return JsonResponse({"error": "Método não permitido"}, status=405)

def buscar_status(request):
    status_list = list(Status.objects.values("id", "nome"))
    return JsonResponse({"status": status_list})

def exportar_historico_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Histórico de Atividades"

    headers = ["ID", "Paciente", "Parceiros", "Produtos", "Data Criação", "Valor Venda", "Valor Repasse", "Status"]
    ws.append(headers)
    
    status_id = request.GET.get("status")
    paciente_nome = request.GET.get("paciente")
    especialidade_id = request.GET.get("especialidade")
    procedimento_nome = request.GET.get("procedimento")
    parceiro_nome = request.GET.get("parceiro")
    data_criacao = request.GET.get("data_criacao")
    data_aprovacao = request.GET.get("data_aprovacao")

    atividades = Orcamento.objects.select_related("status").prefetch_related(
        "solicitacao_orcamento__paciente",
        "orcamento_parceiros__parceiro",
        "orcamento_parceiros__produto"
    )

    if status_id:
        atividades = atividades.filter(status_id=status_id)
    
    if paciente_nome:
        atividades = atividades.filter(solicitacao_orcamento__paciente__nome__icontains=paciente_nome)
    
    if especialidade_id:
        atividades = atividades.filter(solicitacao_orcamento__especialidade_id=especialidade_id)
    
    if procedimento_nome:
        atividades = atividades.filter(orcamento_parceiros__produto__nome__icontains=procedimento_nome)
    
    if parceiro_nome:
        atividades = atividades.filter(orcamento_parceiros__parceiro__nome__icontains=parceiro_nome)
    
    if data_criacao:
        atividades = atividades.filter(data_criacao=data_criacao)
    
    if data_aprovacao:
        atividades = atividades.filter(data_aprovacao=data_aprovacao)

    for atividade in atividades:
        paciente = atividade.solicitacao_orcamento.paciente.nome if atividade.solicitacao_orcamento else "N/A"
        status = atividade.status.nome if atividade.status else "N/A"
        data_criacao = atividade.data_criacao.strftime("%d/%m/%Y") if atividade.data_criacao else "N/A"

        for orcamento_parceiro in atividade.orcamento_parceiros.all():
            parceiro = orcamento_parceiro.parceiro.nome
            produto = orcamento_parceiro.produto.nome
            valor_venda = orcamento_parceiro.valor_venda
            valor_repasse = orcamento_parceiro.valor_repasse
        
            ws.append([atividade.id, paciente, parceiro, produto, data_criacao, valor_venda, valor_repasse, status])

    moeda_format = 'R$ #,##0.00'
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = moeda_format

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="historico_atividades.xlsx"'

    wb.save(response)
    return response

# ======================================================================
from rest_framework import viewsets
from .models import Coluna, Card, Orcamento
from rest_framework.decorators import api_view
from rest_framework.response import Response
from .serializers import ColunaSerializer, CardSerializer, OrcamentoSerializer

class ColumnViewSet(viewsets.ModelViewSet):
    queryset = Coluna.objects.all().order_by('ordem')
    serializer_class = ColunaSerializer

class CardViewSet(viewsets.ModelViewSet):
    queryset = Card.objects.all().order_by('prioridade')
    serializer_class = CardSerializer

# @api_view(['GET'])
# def kanban_board(request):
#     statuses = Status.objects.all().order_by('id')
#     board = []

#     for status in statuses:
#         orcamentos = Orcamento.objects.select_related('solicitacao_orcamento__paciente').filter(status=status).order_by('-data_criacao')
#         board.append({
#             "id": status.id,
#             "title": status.nome,
#             "cards": [
#                 {
#                     "id": orc.id,
#                     "title": f"Orçamento #{orc.id}",
#                     "description": f"R$ {orc.valor_total:.2f}",
#                     "paciente": orc.solicitacao_orcamento.paciente.nome if orc.solicitacao_orcamento else "",
#                     "data_criacao": orc.data_criacao.strftime('%d/%m/%Y') if orc.data_criacao else "",
#                 }
#                 for orc in orcamentos
#             ]
#         })

#     return Response(board)

@api_view(['GET'])
def kanban_board(request):
    colunas = Coluna.objects.order_by('ordem').prefetch_related('cards__orcamento__solicitacao_orcamento__paciente')
    serializer = ColunaSerializer(colunas, many=True)
    return Response(serializer.data)

@api_view(['GET'])
def historico_api(request):
    atividades = Orcamento.objects.all().order_by('-data_criacao')
    serializer = OrcamentoSerializer(atividades, many=True)
    return Response(serializer.data)

@api_view(['PATCH'])
def atualizar_orcamento_status(request, pk):
    orcamento = get_object_or_404(Orcamento, id=pk)
    status_id = request.data.get("status")

    if not status_id:
        return Response({"error": "status obrigatório"}, status=400)

    orcamento.status_id = status_id
    orcamento.save()
    return Response({"success": True})

@api_view(['DELETE'])
def deletar_coluna(request, coluna_id):
    try:
        coluna = Coluna.objects.get(id=coluna_id)
        coluna.delete()
        return Response({'mensagem': 'Coluna excluída com sucesso'}, status=204)
    except Coluna.DoesNotExist:
        return Response({'erro': 'Coluna não encontrada'}, status=404)
    
from django.views.decorators.csrf import ensure_csrf_cookie
from django.http import JsonResponse

@ensure_csrf_cookie
def get_csrf_token(request):
    return JsonResponse({"message": "Token set"})